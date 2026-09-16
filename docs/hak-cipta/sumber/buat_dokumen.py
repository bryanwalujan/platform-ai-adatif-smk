from pathlib import Path
import subprocess, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
REPO=ROOT.parents[1]

def diagram(name,body):
    src='digraph G { graph [bgcolor="white",pad="0.25",ranksep="0.5"]; node [shape=box,style="rounded,filled",fillcolor="#EAF2FA",color="#24618A",fontname="DejaVu Sans",fontsize=12,margin="0.15"]; edge [color="#53758A",fontname="DejaVu Sans",fontsize=10]; '+body+' }'
    p=ROOT/'diagram'/f'{name}.dot'; p.write_text(src)
    for ext in ['png','svg']:
        subprocess.run(['dot',f'-T{ext}',str(p),'-o',str(p.with_suffix('.'+ext))],check=True)

diagram('01-arsitektur','''rankdir=TB;
user [label="Siswa • Guru • Admin sekolah"];
flutter [label="Antarmuka Flutter\nMateri • Kuis • Proyek • Kemajuan"];
api [label="Laravel REST API\nAutentikasi Sanctum"];
school [label="Konteks sekolah dari akun\nValidasi peran dan akses mata pelajaran"];
services [label="Layanan pembelajaran\nRekomendasi adaptif • BKT • Penilaian PBL"];
db [label="Basis data relasional\nAkun, mapel, topik, nilai, riwayat"];
files [label="Penyimpanan privat\nLampiran materi dan proyek"];
user -> flutter -> api -> school -> services; services -> db; school -> files [label="Akses lampiran berizin"];
''')
diagram('02-alur-pbl','''rankdir=TB;
a [label="Guru menetapkan masalah/proyek\nmelalui materi dan arahan pembelajaran"];
b [label="Siswa memilih mapel dan mempelajari materi\nTeks • Gambar • Video • Audio • Dokumen"];
c [label="Siswa mengerjakan kuis\nSistem memperbarui indikator penguasaan"];
d [label="Siswa melihat rekomendasi dan level PBL\nDasar • Menengah • Lanjutan"];
e [label="Siswa menyelidiki masalah dan menyusun hasil\nKegiatan belajar dipandu guru"];
f [label="Siswa mengirim proyek PBL\nJudul, penjelasan dan/atau lampiran"];
g [label="Guru meninjau bukti dan memberi nilai rubrik\nDisertai umpan balik"];
h [label="Siswa membaca hasil dan melakukan refleksi\nMelanjutkan pembelajaran"];
a -> b -> c -> d -> e -> f -> g -> h; h -> b [label="Siklus berikutnya"];
''')
diagram('03-relasi-data','''rankdir=LR;
s [label="schools"];
u [label="users\nPeran dan school_id"];
m [label="subjects"];
t [label="topics"];
mat [label="materials\nmedia_files (JSON)"];
q [label="quizzes"];
qq [label="quiz_questions"];
p [label="pbl_projects\nLampiran, rubrik, nilai"];
st [label="subject_student\nKeanggotaan siswa"];
gt [label="subject_teacher\nPengampu guru"];
master [label="student_topic_mastery"];
s -> u [label="1:N"]; s -> m [label="1:N"]; m -> t [label="1:N"]; t -> mat [label="1:N"]; t -> q [label="1:N"]; q -> qq [label="1:N"];
u -> st; m -> st; u -> gt; m -> gt; u -> p; m -> p; t -> p [label="opsional"]; u -> master; t -> master;
''')
diagram('04-akses-sekolah','''rankdir=TB;
a [label="Permintaan API dan token akun"];
b [label="Autentikasi akun"];
c [label="Ambil sekolah dari akun\nPastikan sekolah aktif"];
d [label="Batasi pencarian data ke school_id akun"];
e [label="Periksa peran, kepemilikan\ndan keanggotaan/pengampu mapel"];
f [label="Kirim data atau lampiran yang diizinkan"];
x [label="Tolak permintaan\nTidak terautentikasi / tidak berhak / tidak ditemukan",fillcolor="#FCEBE8"];
a -> b -> c -> d -> e -> f; b -> x [label="gagal"]; c -> x [label="gagal"]; e -> x [label="gagal"];
''')

def newdoc(subtitle):
    d=Document(); sec=d.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21)
    sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
    st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11); st.paragraph_format.space_after=Pt(7)
    for n in ['Title','Heading 1','Heading 2']: d.styles[n].font.color.rgb=RGBColor.from_string('174B70')
    sec.header.paragraphs[0].text='PIJAR PBL  /  DOKUMENTASI PROGRAM KOMPUTER'
    foot=sec.footer.paragraphs[0]; foot.text=subtitle+'  •  '
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
    d.add_paragraph('PIJAR PBL','Title'); d.add_paragraph(subtitle,'Subtitle')
    d.add_paragraph('Platform Pembelajaran Adaptif Berbasis PBL untuk Siswa SMK','Heading 1')
    d.add_paragraph('Nama pencipta: [ISI NAMA LENGKAP]\nInstitusi: [ISI INSTITUSI]\nPemegang hak cipta: [ISI IDENTITAS]\nTahun penciptaan: [ISI TAHUN]\nTanggal pertama diumumkan: [ISI TANGGAL, JIKA ADA]')
    d.add_paragraph('Edisi dokumentasi: 16 September 2026\nAcuan kode: Laravel 5e3356b • Flutter c8ec620')
    d.add_paragraph('Pijar PBL merupakan nama usulan yang digunakan dalam dokumen ini. Antarmuka pada versi kode acuan masih menggunakan nama Belajar Adaptif. Samakan penamaan sebelum menyiapkan berkas akhir.')
    d.add_paragraph('Identitas dan riwayat penciptaan harus diisi oleh pemilik karya. Paket ini mendokumentasikan aplikasi; tidak menyatakan bahwa pencatatan hak cipta telah diajukan atau disetujui.')
    return d

def page(d,title): d.add_page_break(); d.add_heading(title,level=1)
def para(d,text): d.add_paragraph(text)
def steps(d,items):
    for i,t in enumerate(items,1): d.add_paragraph(f'{i}. {t}')
def table(d,headers,rows):
    t=d.add_table(rows=1, cols=len(headers)); t.style='Light Shading Accent 1'
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)

def pic(d,name,width=6.3): d.add_picture(str(ROOT/'diagram'/f'{name}.png'),width=Inches(width))

manual=newdoc('Panduan Penggunaan Aplikasi')
page(manual,'1. Mengenal aplikasi')
para(manual,'Pijar PBL membantu siswa SMK mempelajari materi, mengerjakan kuis, menyelesaikan proyek PBL, dan menindaklanjuti umpan balik guru. Platform menggabungkan materi multimedia, pemantauan penguasaan topik, dan rekomendasi belajar. Satu instalasi dapat melayani beberapa sekolah dengan akun serta data pembelajaran yang dibatasi per sekolah.')
table(manual,['Peran','Kegiatan utama'],[['Siswa','Mengikuti mapel, mempelajari materi, mengerjakan kuis, mengirim proyek, membaca nilai dan rekomendasi.'],['Guru','Mengelola mapel dan konten, meninjau kemajuan siswa, menilai proyek dengan rubrik.'],['Admin sekolah','Mengelola identitas sekolah, pengguna, persetujuan guru, dan pengampu mapel dalam sekolahnya.'],['Operator server','Menyiapkan layanan, sekolah awal, dan akun admin melalui fasilitas server. Peran operasional ini bukan menu lintas sekolah di Flutter.']])
para(manual,'Persiapan: gunakan versi aplikasi terbaru, koneksi internet, alamat email aktif, dan kode sekolah dari admin. Siswa juga memerlukan kode kelas/mapel dari guru. Pemutaran media bergantung pada format dan kemampuan perangkat.')
para(manual,'Urutan panduan: akun (bagian 2), admin (3), guru dan konten (4–5), siswa (6–7), penilaian (8), serta bantuan (9).')
page(manual,'2. Pendaftaran dan masuk')
steps(manual,['Buka aplikasi dan pilih pendaftaran akun. Isi nama, email aktif, kata sandi minimal delapan karakter, peran siswa/guru, dan kode sekolah.', 'Periksa kode sekolah. Kode sekolah menentukan lingkungan sekolah akun; kode ini berbeda dari kode kelas/mapel.', 'Kirim pendaftaran, buka email, lalu masukkan kode verifikasi enam digit. Gunakan kirim ulang jika diperlukan dan periksa folder spam.', 'Guru menunggu persetujuan admin sekolah sebelum memakai fitur guru. Verifikasi email dan persetujuan guru merupakan dua tahapan terpisah.', 'Untuk masuk berikutnya, gunakan email dan kata sandi. Periksa identitas sekolah pada akun sebelum memulai aktivitas.', 'Gunakan fitur lupa kata sandi pada halaman masuk jika tidak dapat mengakses akun. Ikuti instruksi email dan formulir pemulihan.', 'Keluar dari akun setelah menggunakan perangkat bersama melalui halaman profil.'])
para(manual,'Satu email digunakan untuk satu akun dalam aplikasi. Sekolah ditetapkan dari akun dan tidak dapat dipindah melalui perubahan profil. Hubungi admin jika mendaftar pada sekolah yang keliru.')
page(manual,'3. Panduan admin sekolah')
steps(manual,['Masuk memakai akun admin yang disiapkan operator. Buka Dashboard untuk melihat ringkasan sekolah.', 'Buka Pengaturan Sekolah. Periksa dan simpan identitas sekolah. Bagikan kode sekolah hanya kepada calon pengguna sekolah tersebut.', 'Jika kode perlu diganti, gunakan fasilitas regenerasi kode. Berikan kode baru kepada calon pendaftar.', 'Buka persetujuan guru. Periksa identitas calon guru sebelum menyetujui atau menolak.', 'Buka pengelolaan pengguna untuk memeriksa daftar dan status pengguna sekolah.', 'Buka pengelolaan mata pelajaran untuk memeriksa mapel dan mengatur guru pengampu sesuai kebutuhan.'])
para(manual,'Admin bekerja dalam sekolah akun sendiri. Pembuatan sekolah tambahan dan admin pertamanya dilakukan operator server; tidak tersedia sebagai menu superadmin lintas sekolah pada Flutter.')
para(manual,'Hasil yang diperiksa: nama sekolah benar, guru yang sah sudah disetujui, dan setiap guru terhubung dengan mata pelajaran yang diampunya.')
page(manual,'4. Guru: menyiapkan mata pelajaran')
steps(manual,['Masuk sebagai guru yang sudah disetujui, lalu buka tab Mapel.', 'Buat mata pelajaran melalui fasilitas tambah atau pilih mapel yang sudah diampu. Isi informasi yang diminta pada formulir.', 'Buka detail mapel. Bagikan kode kelas/mapel kepada siswa dari sekolah yang sama atau tambahkan siswa melalui tombol Tambah.', 'Pilih Kelola Konten. Halaman menyediakan bagian Materi, Topik, Kuis, dan Soal.', 'Buat topik terlebih dahulu. Beri judul dan penjelasan yang sesuai dengan tujuan pembelajaran.', 'Gunakan RPP pada detail mapel untuk mengatur rencana pembelajaran per pertemuan sesuai formulir yang tersedia.', 'Tuliskan masalah pemantik, tujuan, instruksi proyek, bukti yang diharapkan, dan kriteria keberhasilan dalam materi/arahan pembelajaran.'])
para(manual,'Contoh rancangan PBL untuk Matematika: siswa merancang denah ruang kerja, menghitung luas dan kebutuhan bahan, kemudian mengirim penjelasan perhitungan beserta gambar denah. Contoh ini merupakan bahan demonstrasi, bukan data siswa nyata.')
para(manual,'PBL tetap menjadi pendekatan pembelajaran. Rencana penyelidikan, diskusi, dan refleksi dapat dipandu guru melalui materi dan forum; dokumen ini tidak mengasumsikan adanya modul khusus penjadwalan proyek atau kerja kelompok otomatis.')
page(manual,'5. Guru: materi multimedia dan kuis')
steps(manual,['Dari Kelola Konten, pilih Materi dan buka formulir materi baru. Pilih topik tujuan.', 'Isi judul materi, perkiraan durasi, serta penjelasan yang membantu siswa memahami konteks. Penjelasan dapat dilengkapi atau digantikan lampiran/tautan; materi tidak boleh sepenuhnya kosong.', 'Pada media pembelajaran, pilih Tambah lampiran. Pilih gambar, video, audio, dokumen, atau berkas sumber yang relevan.', 'Jika menggunakan tautan video, masukkan alamat HTTP/HTTPS yang benar. Tautan berkas video langsung dapat diputar; tautan halaman seperti YouTube dapat disalin untuk dibuka di browser.', 'Tekan Simpan materi untuk siswa dan tunggu proses unggahan selesai. Periksa hasil melalui pratinjau.', 'Untuk memperbarui, pilih Edit pada menu materi. Lampiran lama tetap tersimpan. Gunakan Hapus lampiran ini hanya untuk lampiran yang ingin dibuang; tersedia pembatalan penghapusan sebelum disimpan.', 'Untuk kuis, buka Kuis dan buat kuis pada topik terkait. Selanjutnya buka Soal, pilih topik/kuis, isi pertanyaan, pilihan jawaban, jawaban benar, dan penjelasan.'])
table(manual,['Jenis','Format yang diterima'],[['Gambar','JPG, JPEG, PNG, GIF, WebP'],['Video dan audio','MP4, MOV, WebM, MP3, WAV, OGG, M4A'],['Dokumen dan arsip','PDF, TXT, DOC, DOCX, PPT, PPTX, ZIP']])
para(manual,'Batas per materi: lima lampiran, 50 MB per file, total 100 MB. Server dapat menolak lebih awal apabila batas hosting lebih kecil. Periksa kembali topik sebelum menyimpan agar materi masuk ke mapel yang tepat.')
page(manual,'6. Siswa: mengikuti dan mempelajari materi')
steps(manual,['Masuk sebagai siswa. Buka Mapel dan gunakan fasilitas bergabung dengan kode kelas/mapel dari guru.', 'Pilih mata pelajaran, lalu buka daftar topik dan topik yang ingin dipelajari.', 'Pilih materi. Baca penjelasan dan amati gambar. Untuk video/audio gunakan tombol Putar video atau Dengarkan audio.', 'Gunakan kontrol putar/jeda dan geser posisi media bila didukung. Gunakan Unduh / simpan / bagikan untuk menyimpan dokumen atau berkas yang perlu aplikasi lain.', 'Tekan Selesai belajar setelah mempelajari materi. Durasi dan interaksi belajar digunakan sebagai catatan aktivitas sesuai mekanisme aplikasi.', 'Buka kuis dari topik, jawab soal, dan kirim jawaban. Perhatikan hasil sebagai bahan perbaikan belajar.', 'Gunakan Forum Diskusi Topik Ini untuk mengajukan pertanyaan atau membahas pemecahan masalah dengan pengguna yang memiliki akses topik.'])
para(manual,'Gunakan tab Adaptif dan pilih mapel untuk melihat rekomendasi. Penguasaan topik dan rekomendasi membantu menentukan prioritas belajar; indikator tersebut bukan jaminan penguasaan seluruh kompetensi praktik siswa.')
page(manual,'7. Siswa: mengumpulkan proyek PBL')
steps(manual,['Pelajari masalah dan arahan guru. Siapkan hasil pekerjaan serta bukti proses yang relevan dengan mata pelajaran.', 'Pada detail topik pilih Buat Proyek PBL untuk Topik Ini. Pastikan konteks mapel/topik dan level yang ditampilkan sudah sesuai.', 'Isi Judul proyek. Pada Jawaban / penjelasan karya, jelaskan masalah, cara kerja, hasil, dan kesimpulan.', 'Pilih Tambah lampiran jika membutuhkan gambar, rekaman demonstrasi, narasi audio, laporan, presentasi, atau berkas sumber.', 'Periksa kembali semua berkas. Maksimal lima lampiran, 50 MB per file, dan total 100 MB. Minimal tersedia penjelasan teks atau lampiran.', 'Tekan Kirim proyek ke guru dan tunggu hingga berhasil. Hindari keluar dari halaman selama unggahan berlangsung.', 'Pantau status proyek dan notifikasi. Setelah dinilai, buka hasil/feedback proyek untuk melihat skor, rincian rubrik, dan masukan guru.', 'Gunakan masukan untuk refleksi dan tugas berikutnya. Proyek yang telah dinilai tidak dapat dihapus atau dinilai ulang melalui alur yang tersedia.'])
para(manual,'Pilih media berdasarkan bukti kompetensi, bukan jumlah file. Contohnya: perhitungan tertulis untuk Matematika, rekaman dialog untuk Bahasa, video prosedur untuk mata pelajaran kejuruan. Semua contoh ini adalah skenario demonstrasi.')
page(manual,'8. Guru: menilai dan memantau belajar')
steps(manual,['Buka daftar proyek yang menunggu penilaian dari dashboard guru. Pilih proyek siswa.', 'Periksa identitas siswa, mapel, topik, penjelasan karya, dan seluruh lampiran yang relevan.', 'Beri skor 0–100 untuk setiap kriteria rubrik. Tuliskan catatan kriteria dan umpan balik umum yang menjelaskan kelebihan serta perbaikan.', 'Periksa nilai akhir sebelum menyimpan. Penilaian yang sudah disimpan tidak dapat diulang melalui alur saat ini.', 'Buka daftar siswa dan kemajuannya untuk meninjau hasil belajar. Gunakan Monitor AI Adaptif pada detail mapel untuk melihat penguasaan dan rekomendasi siswa dalam mapel itu.'])
table(manual,['Kriteria','Bobot','Fokus'],[['Pemecahan Masalah & Inisiatif','25%','Penalaran, solusi, inisiatif pendekatan'],['Ketepatan Metode & Hasil','35%','Prosedur, akurasi, bukti'],['Pemahaman Materi','25%','Konsep dan penerapannya'],['Komunikasi & Dokumentasi','15%','Kejelasan penyajian dan dokumentasi']])
para(manual,'Contoh nilai (ilustrasi): skor 80, 90, 85, dan 80 menghasilkan (80 × 0,25) + (90 × 0,35) + (85 × 0,25) + (80 × 0,15) = 84,75. Guru menilai substansi sesuai mapel; multimedia merupakan bentuk bukti, bukan syarat untuk nilai tinggi.')
page(manual,'9. Bantuan dan pemeriksaan hasil')
table(manual,['Kendala','Tindakan'],[['Kode sekolah ditolak','Minta kode sekolah aktif kepada admin; jangan gunakan kode kelas/mapel.'],['Guru belum dapat mengelola konten','Pastikan email terverifikasi dan persetujuan admin sudah selesai.'],['Mapel atau materi tidak terlihat','Periksa akun/sekolah, keanggotaan siswa, pengampu guru, dan topik yang dipilih.'],['Unggahan gagal','Periksa koneksi, format, jumlah file, ukuran, dan batas unggahan hosting.'],['Video/audio tidak diputar','Coba unduh dan buka di aplikasi lain. Gunakan format/codec yang didukung perangkat.'],['Tautan file tidak dapat dibuka dari browser','Akses melalui akun yang berhak di aplikasi; lampiran membutuhkan autentikasi.'],['Nilai/rekomendasi belum terlihat','Muat ulang halaman, pastikan mapel benar dan kegiatan kuis/penilaian sudah tersimpan.']])
para(manual,'Pemeriksaan singkat sesudah pemasangan: guru membuat materi, siswa membukanya, siswa mengirim proyek, guru menilai, dan siswa menerima hasil. Lakukan pada akun uji sekolah yang sama. Pemeriksaan perangkat nyata masih diperlukan untuk pemutar, unduhan, email, serta koneksi server.')
para(manual,'Kontak dukungan sekolah: [ISI NAMA/KONTAK]\nVersi aplikasi yang digunakan: [ISI VERSI]\nAlamat layanan: [ISI ALAMAT PUBLIK JIKA DIPERLUKAN]')
manual.save(ROOT/'01-Panduan-Penggunaan.docx')

art=newdoc('Deskripsi dan Artefak Aplikasi')
page(art,'A. Deskripsi ciptaan')
para(art,'Pijar PBL adalah program komputer berupa platform pembelajaran adaptif berbasis PBL untuk siswa SMK. Aplikasi menyediakan antarmuka Flutter dan layanan backend Laravel untuk mengelola mata pelajaran, topik, materi multimedia, kuis, proyek siswa, penilaian rubrik, diskusi, dan riwayat pembelajaran. Pengguna dikelompokkan sebagai siswa, guru, dan admin sekolah.')
para(art,'Sistem membantu siswa menentukan prioritas belajar berdasarkan indikator penguasaan dan rekomendasi per mata pelajaran. Bukti proyek dapat disampaikan sebagai teks, gambar, video, audio, dokumen, atau arsip. Guru memberi penilaian melalui empat kriteria universal sehingga pembelajaran dapat diterapkan pada berbagai mata pelajaran SMK.')
para(art,'Dukungan beberapa sekolah diwujudkan melalui identitas sekolah pada akun, pembatasan data, validasi relasi, serta kontrol akses pada lampiran. Admin sekolah bekerja pada lingkungan sekolahnya sendiri. Penyimpanan file pembelajaran memerlukan akses terautentikasi.')
para(art,'Komponen adaptif mencakup indikator penguasaan, rekomendasi berbasis aturan dan riwayat, serta Bayesian Knowledge Tracing (BKT), yaitu estimasi peluang penguasaan dari urutan hasil jawaban. Keberadaan BKT tidak berarti seluruh rekomendasi dihasilkan oleh model generatif atau chatbot. Penilaian proyek tetap diberikan guru.')
para(art,'Ruang lingkup karya yang didokumentasikan: implementasi aplikasi, rancangan interaksi, struktur pengolahan pembelajaran, integrasi media, dan integrasi pembatasan sekolah pada versi acuan. Flutter, Laravel, serta pustaka pihak ketiga merupakan dependensi; dokumen ini tidak menyatakan dependensi tersebut sebagai ciptaan pemohon.')
page(art,'B. Inventaris fitur dan bukti kode')
table(art,['Fitur','Lokasi implementasi utama'],[['Akun dan verifikasi','app/Http/Controllers/Api/AuthController.php'],['Identitas dan isolasi sekolah','app/Models/Concerns/BelongsToSchool.php; middleware sekolah'],['Konten multimedia','app/Services/MaterialMediaService.php; MaterialController.php'],['Manajemen konten Flutter','lib/screens/teacher/content_management_screen.dart'],['Editor dan pembaca materi','lib/screens/teacher/material_editor_screen.dart; lib/screens/material/material_detail_screen.dart'],['Pengumpulan dan rubrik PBL','app/Models/PblProject.php; PblProjectController.php; TeacherController.php'],['Adaptasi pembelajaran','app/Services/AdaptiveEngineService.php; BayesianKnowledgeTracingService.php'],['Penampil lampiran','lib/widgets/project_attachments.dart']])
para(art,'Path app/ merujuk proyek Laravel; path lib/ merujuk proyek Flutter. Kode cuplikan disertakan terpisah dengan rentang baris dan checksum. Diagram adalah ilustrasi teknis berdasarkan kode, bukan tangkapan layar antarmuka.')
page(art,'C. Diagram arsitektur')
pic(art,'01-arsitektur',5.8)
para(art,'Gambar C.1. Arsitektur ringkas Flutter–Laravel. Permintaan data melewati autentikasi, konteks sekolah, dan pemeriksaan akses. Diagram menampilkan jalur utama, bukan seluruh komponen deployment.')
page(art,'D. Alur pembelajaran berbasis PBL')
pic(art,'02-alur-pbl',4.6)
para(art,'Gambar D.1. Alur operasional pembelajaran. Penetapan masalah, penyelidikan, dan refleksi merupakan kegiatan pedagogis yang didukung materi, forum, serta proyek; bukan klaim adanya formulir khusus untuk setiap tahap.')
page(art,'E. Relasi data inti')
pic(art,'03-relasi-data',6.4)
para(art,'Gambar E.1. Relasi konseptual inti (subset skema). Panah menunjukkan arah hubungan induk ke data terkait. subject_student dan subject_teacher menghubungkan akun dengan mapel. Tabel aktivitas, notifikasi, percobaan kuis, RPP, dan parameter BKT tidak digambar agar bagan tetap terbaca. Pembatas sekolah juga berlaku pada model terkait yang membawa school_id.')
page(art,'F. Alur pembatasan akses')
pic(art,'04-akses-sekolah',5.8)
para(art,'Gambar F.1. Konteks sekolah diturunkan dari akun terautentikasi. Mengubah ID sekolah pada request tidak memberikan hak akses ke sekolah lain. Ilustrasi ini menjelaskan kontrol yang diterapkan, bukan klaim bahwa sistem bebas dari seluruh kemungkinan kerentanan.')
page(art,'G. Skenario demonstrasi dan tangkapan layar')
para(art,'Gunakan akun dan data demonstrasi. Tangkapan layar aplikasi nyata belum disertakan dalam paket ini. Ambil gambar berikut pada versi terbaru, lalu tempelkan pada lembar tangkapan layar terpisah. Jangan menampilkan kata sandi, token, kode verifikasi, atau identitas siswa nyata.')
table(art,['ID','Tampilan','Bukti yang ditunjukkan'],[['SS-01','Login/pendaftaran','Peran dan kode sekolah'],['SS-02','Dashboard guru + identitas sekolah','Lingkungan sekolah dan fungsi guru'],['SS-03','Detail mapel','Kode kelas, konten, RPP, siswa'],['SS-04','Kelola Konten','Daftar materi dan filter topik'],['SS-05','Editor materi','Lampiran gambar, audio/video, dokumen'],['SS-06','Materi pada akun siswa','Gambar dan kontrol media'],['SS-07','Kuis dan hasil','Evaluasi pembelajaran'],['SS-08','Rekomendasi adaptif per mapel','Prioritas topik dan level PBL'],['SS-09','Pengiriman proyek PBL','Judul, penjelasan, lampiran'],['SS-10','Penilaian guru','Empat kriteria dan nilai berbobot'],['SS-11','Hasil proyek siswa','Nilai dan feedback'],['SS-12','Pengaturan sekolah admin','Identitas dan pengelolaan sekolah']])
page(art,'H. Validasi dan identitas versi')
para(art,'Versi acuan yang telah diperiksa pada pekerjaan implementasi sebelumnya: Laravel 5e3356b dan Flutter c8ec620. Catatan pengujian: 62 tes backend (299 assertions) lulus; tujuh tes Flutter terkait materi, PBL, dan sekolah lulus; analisis empat file Flutter terkait tidak menemukan masalah. Ini adalah ringkasan hasil sebelumnya, bukan klaim pengujian ulang saat penyusunan dokumen.')
para(art,'Batas bukti: keseluruhan suite Flutter masih memiliki kegagalan pada tes template penghitung lama di widget_test.dart. Pengujian otomatis tidak membuktikan pemutaran semua codec pada semua perangkat, pengiriman email hosting, maupun keberhasilan deployment produksi. Tidak ada bukti uji pengguna atau hasil penelitian baru yang dibuat untuk dokumen ini.')
para(art,'Metadata berkas dan SHA-256 tersedia dalam MANIFEST.json. Checksum membantu memeriksa bahwa berkas paket tidak berubah; checksum bukan bukti mandiri kepemilikan atau tanggal penciptaan.')
para(art,'Sebelum digunakan: isi identitas pencipta/pemegang hak dan riwayat penciptaan berdasarkan data sebenarnya, tetapkan nama final, lengkapi tangkapan layar, lalu tinjau dokumen hasil ekspor. Nomor permohonan, tanda tangan, dan pengesahan tidak dibuat dalam paket ini.')
art.save(ROOT/'02-Deskripsi-dan-Artefak.docx')

shots=newdoc('Lembar Tangkapan Layar')
for i,(title,caption) in enumerate([('Akun dan sekolah','Tampilan pendaftaran dengan pilihan peran dan kode sekolah.'),('Dashboard guru','Ringkasan guru dan identitas sekolah aktif.'),('Detail mapel','Pengelolaan konten dan siswa dalam mata pelajaran.'),('Manajemen konten','Daftar materi serta filter topik.'),('Editor materi multimedia','Form materi dengan beberapa jenis lampiran.'),('Materi siswa','Pembacaan materi dan pemutaran media.'),('Kuis dan hasil','Proses evaluasi penguasaan topik.'),('Rekomendasi adaptif','Rekomendasi sesuai mata pelajaran.'),('Pengumpulan PBL','Pengumpulan penjelasan dan bukti multimedia.'),('Rubrik penilaian','Empat kriteria penilaian proyek oleh guru.'),('Feedback siswa','Nilai dan umpan balik yang diterima siswa.'),('Pengaturan sekolah','Identitas sekolah pada akun admin.')],1):
    page(shots,f'SS-{i:02d}. {title}'); para(shots,'[TEMPELKAN TANGKAPAN LAYAR ASLI DI SINI]')
    para(shots,'Status: belum dilengkapi. Gambar tidak dibuat atau disimulasikan oleh penyusun dokumen.')
    para(shots,f'Keterangan: {caption}\nTanggal pengambilan: [ISI]\nVersi aplikasi: [ISI]\nPerangkat: [ISI]\nPeran akun uji: [ISI]')
shots.save(ROOT/'03-Template-Tangkapan-Layar.docx')

code=newdoc('Cuplikan Kode Implementasi')
selections=[('Isolasi model per sekolah','app/Models/Concerns/BelongsToSchool.php',1,85),('Penyimpanan materi multimedia','app/Services/MaterialMediaService.php',1,120),('Penilaian rubrik PBL','app/Models/PblProject.php',73,180),('Rekomendasi dan level proyek','app/Services/AdaptiveEngineService.php',237,275)]
index=[]
for title,path,start,end in selections:
    full=REPO/path; lines=full.read_text().splitlines(); end=min(end,len(lines)); excerpt='\n'.join(f'{n:04d}  {lines[n-1]}' for n in range(start,end+1))
    page(code,title); para(code,f'Laravel {path}, baris {start}–{end}. Versi 5e3356b. Cuplikan parsial; dependensi dan konteks lengkap terdapat pada repositori.')
    p=code.add_paragraph(); r=p.add_run(excerpt); r.font.name='DejaVu Sans Mono'; r.font.size=Pt(7)
    p.paragraph_format.space_after=Pt(0)
    index.append({'file':path,'start':start,'end':end,'sha256_full_source':hashlib.sha256(full.read_bytes()).hexdigest()})
(ROOT/'sumber'/'indeks-cuplikan.json').write_text(json.dumps(index,indent=2))
code.save(ROOT/'04-Cuplikan-Kode.docx')
print(ROOT)
